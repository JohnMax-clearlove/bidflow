"""本地资料导入、可追溯解析、中文检索及业绩台账导入。"""

from __future__ import annotations

import csv
import importlib.metadata
import importlib.util
import io
import json
import os
import re
import shutil
import sqlite3
import tempfile
from contextlib import closing
from datetime import date, datetime
from pathlib import Path

from .utils import atomic_json, json_hash, sha256_file, utc_now
from .parser_helpers import parse_document

PARSER_VERSION = "bidflow-parse-0.2.0"
CATEGORY_DIRS = {
    "招标文件": "01_招标文件", "澄清补遗": "01_招标文件",
    "公司证照": "02_公司证照", "公司资质": "03_公司资质",
    "荣誉奖项": "04_荣誉奖项", "人员证书": "05_人员证书",
    "企业业绩证明": "06_企业业绩证明", "人员业绩证明": "07_人员业绩证明",
    "其他补充资料": "08_其他补充资料", "方法论": "08_其他补充资料",
    "历史章节": "08_其他补充资料", "业绩台账": "08_其他补充资料",
}
DOCUMENT_EXTENSIONS = {".md", ".txt", ".docx", ".pdf", ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}


def _resolve_source(project, source) -> Path:
    source = Path(source).expanduser()
    if not source.is_absolute():
        candidates = [project.root / source]
        library = project.meta.get("company_library")
        if library:
            library_path = Path(library)
            if not library_path.is_absolute():
                library_path = project.root / library_path
            candidates.append(library_path / source)
        source = next((p for p in candidates if p.is_file()), candidates[0])
    source = source.resolve()
    if not source.is_file():
        raise ValueError(f"未找到本地文件：{source}")
    return source


def _copy_source(project, source: Path, category: str) -> tuple[Path, str]:
    if category not in CATEGORY_DIRS:
        raise ValueError(f"不支持的资料类别：{category}")
    digest = sha256_file(source)
    relative = Path("01_输入文件") / CATEGORY_DIRS[category] / source.name
    target = project.safe_path(relative)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists() and sha256_file(target) != digest:
        relative = relative.with_name(f"{source.stem}_{digest[:12]}{source.suffix}")
        target = project.safe_path(relative)
    if target.exists():
        if sha256_file(target) != digest:
            raise ValueError("目标名称冲突，已保留原件，请为输入文件更名后重试。")
    else:
        # 排他创建避免覆盖已有原件；复制前后核验可以发现复制过程中源文件变化。
        try:
            with source.open("rb") as incoming, target.open("xb") as outgoing:
                shutil.copyfileobj(incoming, outgoing)
        except FileExistsError:
            if sha256_file(target) != digest:
                raise ValueError("导入时出现同名文件，已停止并保留原件。") from None
        if sha256_file(target) != digest or sha256_file(source) != digest:
            raise ValueError("复制期间源文件发生变化，请检查已保留的副本后重新导入。")
    return target, digest


def _ocr_config(project, requested) -> dict:
    settings = project.load("settings", {})
    custom = requested if isinstance(requested, dict) else {}
    configured_paths = custom.get("model_paths", settings.get("ocr_model_paths", {}))
    if bool(custom.get("enabled", True)) if isinstance(requested, dict) else bool(requested):
        if not configured_paths:
            spec = importlib.util.find_spec("rapidocr")
            model_root = Path(spec.origin).parent / "models" if spec and spec.origin else None
            candidates = {
                "det": model_root / "PP-OCRv6_det_small.onnx" if model_root else Path(""),
                "cls": model_root / "ch_ppocr_mobile_v2.0_cls_mobile.onnx" if model_root else Path(""),
                "rec": model_root / "PP-OCRv6_rec_small.onnx" if model_root else Path(""),
            }
            if all(path.is_file() for path in candidates.values()):
                # RapidOCR wheel 中随包安装的本地模型可以直接使用，不触发联网下载。
                configured_paths = {key: str(path) for key, path in candidates.items()}
    config = {
        "enabled": bool(custom.get("enabled", True)) if isinstance(requested, dict) else bool(requested), "dpi": int(custom.get("dpi", 160)),
        "confidence": float(custom.get("confidence", 0.9)),
        "model_paths": configured_paths,
        "docx_render": bool(settings.get("docx_render", True)),
    }
    if not 72 <= config["dpi"] <= 400 or not 0 <= config["confidence"] <= 1:
        raise ValueError("OCR 分辨率需为 72—400 DPI，置信阈值需为 0—1。")
    if not isinstance(config["model_paths"], dict):
        raise ValueError("OCR 模型路径必须是 det、cls、rec 对应路径的字典。")
    # 模型位置按项目解释，迁移整个项目后仍可引用项目内模型。
    config["model_paths"] = {
        str(key): str((project.root / value).resolve() if not Path(value).is_absolute() else Path(value).resolve())
        for key, value in config["model_paths"].items()
    }
    config["model_hashes"] = {
        key: sha256_file(Path(value)) if Path(value).is_file() else "missing"
        for key, value in config["model_paths"].items()
    }
    try:
        config["engine_version"] = importlib.metadata.version("rapidocr")
    except importlib.metadata.PackageNotFoundError:
        config["engine_version"] = "unavailable"
    return config


def _parser_dependencies() -> dict:
    versions = {}
    for package in ("python-docx", "pdfplumber", "pypdfium2", "Pillow", "rapidocr", "onnxruntime"):
        try:
            versions[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            versions[package] = "unavailable"
    pandoc = shutil.which("pandoc")
    versions["pandoc"] = {"path": pandoc, "mtime": Path(pandoc).stat().st_mtime_ns} if pandoc else "unavailable"
    return versions


def _upsert_file(project, source: Path, target: Path, digest: str, category: str, **extra) -> tuple[list, dict]:
    files = project.load("files")
    identifier = "DOC" + digest[:12]
    previous = next((f for f in files if f["id"] == identifier), None)
    if previous and previous["sha256"] != digest:
        raise ValueError("文件编号哈希前缀冲突，请联系维护人员处理。")
    record = dict(previous or {})
    record.update({
        "id": identifier, "path": previous["path"] if previous else target.relative_to(project.root).as_posix(),
        "sha256": digest, "category": previous["category"] if previous else category,
        "original_name": previous["original_name"] if previous else source.name, "source": previous["source"] if previous else str(source),
        "parser_version": PARSER_VERSION, "parse_status": "pending", "warnings": [],
        "categories": sorted(set((previous or {}).get("categories", []) + [category])),
        "sources": sorted(set((previous or {}).get("sources", []) + [str(source)])),
        **extra,
    })
    files = [f for f in files if f["id"] != identifier] + [record]
    return files, record


def ingest(project, source, category="招标文件", ocr=False) -> dict:
    """导入原件并生成可重建解析结果；语义规则提取另交 Agent。"""
    source = _resolve_source(project, source)
    if source.suffix.lower() not in DOCUMENT_EXTENSIONS:
        raise ValueError("文档导入支持 DOCX、PDF、Markdown、TXT 和图片；Excel/CSV 请使用台账导入。")
    config = _ocr_config(project, ocr)
    target, digest = _copy_source(project, source, category)
    dependencies = _parser_dependencies()
    cache_key = json_hash({"sha256": digest, "parser_version": PARSER_VERSION, "ocr": config, "dependencies": dependencies})
    cache = project.cache_dir / "parse" / cache_key
    cache.mkdir(parents=True, exist_ok=True)
    structure_path = cache / "structure.json"
    cached = False
    if structure_path.is_file():
        try:
            parsed = json.loads(structure_path.read_text(encoding="utf-8"))
            if parsed.get("cache_key") == cache_key and parsed.get("sha256") == digest:
                cached = (isinstance(parsed.get("blocks"), list) and bool(parsed["blocks"])
                          and all(isinstance(b, dict) and {"text", "kind", "page", "locator", "quality"} <= b.keys() for b in parsed["blocks"])
                          and isinstance(parsed.get("warnings"), list) and parsed.get("status") in {"parsed", "needs_review"}
                          and all(project.safe_path(p).is_file() for p in parsed.get("previews", []))
                          and (not parsed.get("rendered_pdf") or project.safe_path(parsed["rendered_pdf"]).is_file()))
        except (ValueError, OSError, TypeError, KeyError):
            cached = False
    if not cached:
        parsed = parse_document(target, cache, config)
        for block in parsed["blocks"]:
            if block.get("preview"):
                block["preview"] = Path(block["preview"]).relative_to(project.root).as_posix()
        parsed["previews"] = [Path(p).relative_to(project.root).as_posix() for p in parsed.get("previews", [])]
        if parsed.get("rendered_pdf"):
            parsed["rendered_pdf"] = Path(parsed["rendered_pdf"]).relative_to(project.root).as_posix()
        parsed.update({"cache_key": cache_key, "sha256": digest, "parser_version": PARSER_VERSION, "ocr": config, "dependencies": dependencies})
    files, record = _upsert_file(project, source, target, digest, category)
    file_id = record["id"]
    blocks = []
    for index, block in enumerate(parsed["blocks"], 1):
        blocks.append({**block, "id": f"{file_id}-B{index:05d}", "file_id": file_id})
    parsed["blocks"] = blocks
    atomic_json(structure_path, parsed)
    # 清洗稿必须带可回查的块编号。Pandoc 另存参考稿，避免其改写破坏来源映射。
    clean_text = "\n\n".join(
        f"<!-- {b['id']} | {b['locator']} -->\n{b['text']}" for b in blocks if not b.get("excluded_from_clean")
    )
    (cache / "clean.md").write_text(clean_text + "\n", encoding="utf-8")
    if parsed.get("semantic_markdown"):
        (cache / "pandoc.md").write_text(parsed["semantic_markdown"], encoding="utf-8")
    record.update({
        "cache_key": cache_key, "parse_status": parsed["status"], "warnings": parsed["warnings"],
        "page_count": parsed.get("page_count"), "block_count": len(blocks),
        "parsed_at": record.get("parsed_at", utc_now()) if cached and record.get("cache_key") == cache_key else utc_now(), "cache_path": cache.relative_to(project.root).as_posix(),
        "ocr_config": config,
        "page_basis": parsed.get("page_basis", "physical" if parsed.get("page_count") else "unpaginated"),
        "rendered_pdf": parsed.get("rendered_pdf"), "dependencies": dependencies,
    })
    previous_blocks = [b for b in project.load("blocks") if b["file_id"] != file_id]
    project.commit({"files": files, "blocks": previous_blocks + blocks}, reason=f"导入资料：{source.name}")
    reindex(project)
    project.refresh_status()
    return {"status": parsed["status"], "file_id": file_id, "path": record["path"],
            "blocks": len(blocks), "page_count": parsed.get("page_count"), "cache_hit": cached,
            "cache_key": cache_key, "clean_path": str(cache / "clean.md"),
            "structure_path": str(structure_path), "warnings": parsed["warnings"],
            "review_blocks": [b["id"] for b in blocks if b["quality"] == "review"]}


def _index_fingerprint(project) -> str:
    return json_hash({"files": project.load("files"), "blocks": project.load("blocks"),
                      "staff": project.load("staff"), "history": project.load("history")})


def reindex(project) -> dict:
    """数据库只有派生数据，删除后可由主记录完整恢复。"""
    project.cache_dir.mkdir(parents=True, exist_ok=True)
    database = project.cache_dir / "search.sqlite3"
    fd, temporary = tempfile.mkstemp(prefix="search-", suffix=".sqlite3", dir=project.cache_dir)
    os.close(fd)
    temp = Path(temporary)
    records = []
    files = {f["id"]: f for f in project.load("files")}
    for block in project.load("blocks"):
        source = files.get(block["file_id"], {})
        records.append((block["id"], "block", block.get("text", ""), json.dumps({**block, "category": source.get("category"), "categories": source.get("categories", []), "path": source.get("path")}, ensure_ascii=False)))
    for kind in ("staff", "history"):
        for row in project.load(kind):
            content = json.dumps(row, ensure_ascii=False)
            records.append((row["id"], kind, content, content))
    tokenizer = "trigram"
    try:
        with closing(sqlite3.connect(temp)) as connection:
            connection.execute("CREATE TABLE entries (id TEXT, kind TEXT, text TEXT, payload TEXT)")
            connection.executemany("INSERT INTO entries VALUES (?, ?, ?, ?)", records)
            try:
                connection.execute("CREATE VIRTUAL TABLE search_fts USING fts5(text, content='entries', content_rowid='rowid', tokenize='trigram')")
                connection.execute("INSERT INTO search_fts(search_fts) VALUES ('rebuild')")
            except sqlite3.OperationalError:
                tokenizer = "substring"
            connection.execute("CREATE TABLE metadata (name TEXT PRIMARY KEY, value TEXT)")
            connection.executemany("INSERT INTO metadata VALUES (?, ?)", [("fingerprint", _index_fingerprint(project)), ("tokenizer", tokenizer)])
            connection.commit()
        temp.replace(database)
    finally:
        if temp.exists():
            temp.unlink()
    return {"status": "ok", "database": str(database), "entries": len(records), "tokenizer": tokenizer}


def search(project, query, categories=None, limit=20) -> dict:
    query = str(query).strip()
    if not query:
        raise ValueError("请输入检索词。")
    limit = int(limit)
    if not 1 <= limit <= 500:
        raise ValueError("检索结果数量需为 1—500。")
    database = project.cache_dir / "search.sqlite3"
    stale = True
    if database.is_file():
        try:
            with closing(sqlite3.connect(database)) as connection:
                row = connection.execute("SELECT value FROM metadata WHERE name='fingerprint'").fetchone()
                stale = not row or row[0] != _index_fingerprint(project)
        except sqlite3.DatabaseError:
            pass
    if stale:
        reindex(project)
    with closing(sqlite3.connect(database)) as connection:
        tokenizer = connection.execute("SELECT value FROM metadata WHERE name='tokenizer'").fetchone()[0]
        if len(query) >= 3 and tokenizer == "trigram":
            phrase = '"' + query.replace('"', '""') + '"'
            rows = connection.execute("SELECT e.kind, e.payload FROM search_fts f JOIN entries e ON e.rowid=f.rowid WHERE search_fts MATCH ? ORDER BY rank", (phrase,)).fetchall()
            method = "fts5_trigram"
        else:
            rows = connection.execute("SELECT kind, payload FROM entries WHERE instr(lower(text), lower(?)) > 0 ORDER BY rowid", (query,)).fetchall()
            method = "substring"
    selected = set(categories or [])
    if isinstance(categories, str):
        selected = {categories}
    hits = []
    for kind, encoded in rows:
        row = json.loads(encoded)
        actual = set(row.get("categories", []) + ([row["category"]] if row.get("category") else []))
        if selected and not selected.intersection(actual):
            continue
        hits.append({"type": kind, **row})
        if len(hits) >= limit:
            break
    return {"status": "ok", "query": query, "method": method, "count": len(hits), "results": hits}


LEDGER_COLUMNS = {
    "project_name": ["项目名称", "项目名"], "contract_no": ["合同编号", "合同号"],
    "contract_date": ["合同签订时间", "合同签订日期", "签订日期", "合同日期"],
    "client": ["委托单位", "客户名称", "建设单位"], "project_type": ["项目类型"],
    "service_content": ["服务内容", "服务范围"], "investment": ["项目投资额", "投资额"],
    "status": ["项目状态"], "staff_id": ["人员ID", "人员编号", "工号"],
    "staff_name": ["人员姓名", "姓名"], "role": ["人员角色", "担任角色", "岗位"],
    "project_lead": ["项目负责人"], "project_lead_id": ["项目负责人工号", "项目负责人ID"],
    "project_director": ["项目总监"], "project_director_id": ["项目总监工号", "项目总监ID"],
    "members": ["项目成员"], "member_ids": ["项目成员工号", "项目成员ID"],
}
IMPORTANT_COLUMNS = ["project_name", "contract_no", "contract_date", "project_type", "service_content", "investment"]


def _cell_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return str(value).strip()


def _ledger_sheets(path: Path) -> list[tuple[str, list[list[str]], list[str]]]:
    if path.suffix.lower() in {".csv", ".tsv"}:
        raw = path.read_bytes()
        try:
            decoded = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                decoded = raw.decode("gb18030")
            except UnicodeDecodeError as error:
                raise ValueError("台账编码无法识别，请另存为 UTF-8 CSV。") from error
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        return [("CSV", list(csv.reader(io.StringIO(decoded), delimiter=delimiter)), [])]
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("台账支持 XLSX、XLSM、CSV、TSV；旧版 XLS 请先另存为 XLSX。")
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    result = []
    try:
        for sheet in workbook:
            rows, warnings = [], []
            for cells in sheet.iter_rows():
                row = []
                for cell in cells:
                    if cell.data_type == "f":
                        warnings.append(f"{sheet.title}!{cell.coordinate} 是公式，未将公式或缓存值视为事实，请人工提供确认值。")
                        row.append("")
                    else:
                        row.append(_cell_value(cell.value))
                rows.append(row)
            result.append((sheet.title, rows, warnings))
    finally:
        workbook.close()
    return result


def _column_map(headers: list[str], mapping: dict) -> tuple[dict, list[str]]:
    if len(headers) != len(set(headers)):
        # 空列名在导出表中常见，仅拒绝实际重复列名，避免错列。
        names = [h for h in headers if h]
        if len(names) != len(set(names)):
            raise ValueError("台账存在重复列名，请先为重复列设置不同名称。")
    resolved = {}
    for field, aliases in LEDGER_COLUMNS.items():
        requested = mapping.get(field)
        choices = [requested] if requested else [field, *aliases]
        matches = [h for h in headers if h in choices]
        if len(matches) > 1:
            raise ValueError(f"字段 {field} 对应多个列，请通过 mapping 指定唯一列名。")
        if requested and not matches:
            raise ValueError(f"映射指定的列不存在：{requested}")
        if matches:
            resolved[field] = headers.index(matches[0])
    unknown = set(mapping).difference(LEDGER_COLUMNS)
    if unknown:
        raise ValueError(f"不支持的台账字段：{', '.join(sorted(unknown))}")
    return resolved, [f for f in IMPORTANT_COLUMNS if f not in resolved]


def _parts(value: str) -> list[str]:
    return [p.strip() for p in re.split(r"[,，;；、\n]+", value) if p.strip()]


def import_ledger(project, path, mapping=None) -> dict:
    """把管理员台账作为候选事实导入，不生成已核验证据或得分。"""
    source = _resolve_source(project, path)
    sheets = _ledger_sheets(source)
    mapping = mapping or {}
    if not isinstance(mapping, dict):
        raise ValueError("mapping 必须是“标准字段：原始列名”的字典。")
    prepared, warnings, reports = [], [], []
    for sheet, rows, sheet_warnings in sheets:
        warnings.extend(sheet_warnings)
        if not rows or not any(any(v for v in row) for row in rows):
            reports.append({"sheet": sheet, "status": "empty", "missing_columns": IMPORTANT_COLUMNS})
            continue
        # 标题行默认是首个非空行，报告原始行号，不猜测多行表头。
        header_index = next(i for i, row in enumerate(rows) if any(row))
        headers = [_cell_value(v) for v in rows[header_index]]
        columns, missing = _column_map(headers, mapping)
        if "project_name" not in columns and "staff_name" not in columns:
            reports.append({"sheet": sheet, "status": "needs_mapping", "headers": headers, "missing_columns": missing})
            warnings.append(f"工作表“{sheet}”未找到项目名称或人员姓名列，已保留，等待列映射。")
            continue
        reports.append({"sheet": sheet, "status": "ready", "headers": headers,
                        "mapping": {field: headers[i] for field, i in columns.items()}, "missing_columns": missing})
        if missing:
            warnings.append(f"工作表“{sheet}”缺少列：{', '.join(missing)}；相关条件保持材料不足。")
        if not set(columns).intersection({"role", "project_lead", "project_director", "members"}):
            warnings.append(f"工作表“{sheet}”缺少人员角色列，不推断负责人或主持关系。")
        for index in range(header_index + 1, len(rows)):
            row = [_cell_value(v) for v in rows[index]]
            if not any(row):
                continue
            values = {field: row[i] if i < len(row) else "" for field, i in columns.items()}
            raw = {header: row[i] if i < len(row) else "" for i, header in enumerate(headers) if header}
            prepared.append((sheet, index + 1, values, raw))
    target, digest = _copy_source(project, source, "业绩台账")
    files, file_record = _upsert_file(project, source, target, digest, "业绩台账", parser_version="bidflow-ledger-0.2.0")
    file_id = file_record["id"]
    staff = {item["id"]: item for item in project.load("staff")}
    history = {item["id"]: item for item in project.load("history")}
    if file_record.get("ledger_mapping", {}) != mapping and any(
        any(row.get("file_id") == file_id for row in item.get("source_rows", []))
        for item in [*staff.values(), *history.values()]
    ):
        raise ValueError("此台账已按其他列映射导入；请在来源台账中修正列名并生成新版本，避免旧人员关系与新映射混合。")
    added_staff, added_projects, missing_ids = set(), set(), 0
    for sheet, rownum, values, raw in prepared:
        source_row = {"file_id": file_id, "sheet": sheet, "row": rownum, "raw": raw}
        roles = []
        person_values = []
        if values.get("staff_name"):
            person_values.append((values["staff_name"], values.get("staff_id", ""), values.get("role", ""), "staff_name"))
        for field, identity, role in (("project_lead", "project_lead_id", "项目负责人"), ("project_director", "project_director_id", "项目总监"), ("members", "member_ids", "项目成员")):
            names, ids = _parts(values.get(field, "")), _parts(values.get(identity, ""))
            if ids and len(names) != len(ids):
                warnings.append(f"{sheet} 第 {rownum} 行 {field} 姓名与编号数量不一致，编号匹配待核验。")
                ids = []
            person_values.extend((name, ids[i] if ids else "", role, f"{field}:{i}") for i, name in enumerate(names))
        for name, identity, role, field in person_values:
            sid = "STAFF-" + (identity if identity else json_hash([file_id, sheet, rownum, field])[:16])
            if not identity:
                missing_ids += 1
            if sid in staff and staff[sid]["name"] != name:
                warnings.append(f"{sheet} 第 {rownum} 行工号 {identity} 对应不同姓名，分别保留并等待身份核验。")
                staff[sid]["facts"]["identity_status"] = "ambiguous"
                sid += "-" + json_hash([name, file_id, rownum])[:8]
            if sid not in staff:
                staff[sid] = {"id": sid, "name": name, "available": None,
                              "facts": {"identity_status": "recorded_id" if identity else "unresolved", "ledger_staff_id": identity}, "source_rows": []}
                added_staff.add(sid)
            if identity and any(person.get("facts", {}).get("ledger_staff_id") == identity and person["name"] != name for person in staff.values()):
                staff[sid]["facts"]["identity_status"] = "ambiguous"
            if source_row not in staff[sid]["source_rows"]:
                staff[sid]["source_rows"].append(source_row)
            if role:
                roles.append({"staff_id": sid, "name": name, "role": role})
        if not values.get("project_name"):
            continue
        key = ["contract", values["contract_no"]] if values.get("contract_no") else [file_id, sheet, rownum]
        hid = "HIST-" + json_hash(key)[:16]
        facts = {key: value for key, value in values.items() if key in {"client", "project_type", "service_content", "investment", "status"} and value}
        facts["verification"] = "ledger_only"
        if hid not in history:
            history[hid] = {"id": hid, "name": values["project_name"], "contract_no": values.get("contract_no", ""),
                            "contract_date": values.get("contract_date", ""), "roles": [], "facts": facts, "source_rows": []}
            added_projects.add(hid)
        item = history[hid]
        comparison = {"name": values["project_name"], "contract_date": values.get("contract_date", "")}
        conflicts = {key: {"existing": item.get(key), "incoming": value} for key, value in comparison.items() if value and item.get(key) and item[key] != value}
        conflicts.update({key: {"existing": item["facts"].get(key), "incoming": value} for key, value in facts.items() if item["facts"].get(key) and item["facts"][key] != value})
        if conflicts:
            warnings.append(f"{sheet} 第 {rownum} 行合同号 {item['contract_no']} 存在不同记载，未静默覆盖旧值。")
            conflict = {"source_row": source_row, "values": conflicts}
            if conflict not in item.setdefault("conflicts", []):
                item["conflicts"].append(conflict)
        for key, value in comparison.items():
            if not item.get(key):
                item[key] = value
        for key, value in facts.items():
            item["facts"].setdefault(key, value)
        item["roles"].extend(r for r in roles if r not in item["roles"])
        if source_row not in item["source_rows"]:
            item["source_rows"].append(source_row)
    if missing_ids:
        warnings.append(f"{missing_ids} 个人员记载没有稳定人员编号，按来源行单独建候选身份；未按姓名自动合并。")
    warnings = list(dict.fromkeys(warnings))
    file_record.update({"parse_status": "needs_review" if warnings else "parsed", "warnings": warnings,
                        "ledger_mapping": mapping, "ledger_sheets": reports, "ledger_rows": len(prepared)})
    project.commit({"files": files, "staff": list(staff.values()), "history": list(history.values())}, reason=f"导入候选业绩台账：{source.name}")
    reindex(project)
    report = {"status": file_record["parse_status"], "file_id": file_id, "path": file_record["path"],
              "rows": len(prepared), "added_staff": len(added_staff), "added_projects": len(added_projects),
              "sheets": reports, "warnings": warnings}
    report_path = project.safe_path(Path("03_资料匹配") / f"台账导入报告_{file_id}.json")
    report_path.parent.mkdir(parents=True, exist_ok=True)
    atomic_json(report_path, report)
    project.refresh_status()
    return report
