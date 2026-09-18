"""本地 Word 组卷、实际分页及 PDF 链接校验。"""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile
import json
import os
import re
import shutil
import subprocess

from .utils import atomic_json, json_hash, sha256_file, utc_now, write_text

INPUT_COLLECTIONS = ("rules", "evidence", "responses", "sections", "forms", "facts", "settings", "plans", "reviews", "files")

# manual_checks、confirmations、reviews之外的最终审核簿记不进入组卷输入指纹，
# 避免“先确认组卷、再补人工检查”导致确认自我失效形成循环。


def assembly_fingerprint(project) -> str:
    return project.fingerprint(INPUT_COLLECTIONS)


def visual_fingerprint(project) -> str:
    assembly = project.load("assembly", {})
    outputs = []
    for output in assembly.get("outputs", []):
        row = {"volume": output["volume"]}
        for kind in ("docx", "pdf"):
            path = project.safe_path(output[kind]) if output.get(kind) else None
            row[kind] = sha256_file(path) if path and path.is_file() else None
        outputs.append(row)
    positions = [{key: value for key, value in position.items() if key != "output_path"} for position in project.load("positions", [])]
    return json_hash({"outputs": outputs, "positions": positions})


def _confirmed(project, scope, fingerprint) -> bool:
    return any(
        item.get("scope") == scope
        and item.get("fingerprint") == fingerprint
        and item.get("actor")
        and item.get("actor_kind") == "human"
        and item.get("attestation") == "human"
        for item in project.load("confirmations", [])
    )


def _integrity_errors(project):
    errors = []
    for section in project.load("sections", []):
        path = project.safe_path(section["path"])
        if not path.is_file() or not section.get("sha256") or sha256_file(path) != section["sha256"]:
            errors.append(f"章节 {section['id']} 的正文与确认版本不一致")
    for form in project.load("forms", []):
        for field, digest in (("template_path", "template_sha256"), ("output_path", "output_sha256")):
            path = project.safe_path(form[field]) if form.get(field) else None
            if not path or not path.is_file() or not form.get(digest) or sha256_file(path) != form[digest]:
                errors.append(f"表单 {form['id']} 的模板或成稿与记录版本不一致")
    return errors


def _bookmark(target: str) -> str:
    return "BF_" + re.sub(r"[^A-Za-z0-9_]", "_", target)[:22] + "_" + json_hash(target)[:8]


def _anchor(paragraph, name, sequence):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(sequence))
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(sequence))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _field(paragraph, code, text="待更新"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), code)
    run, value = OxmlElement("w:r"), OxmlElement("w:t")
    value.text = text
    run.append(value)
    field.append(run)
    paragraph._p.append(field)


def _link(paragraph, anchor, label):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _inline(paragraph, text):
    # 仅转换文内强调；Markdown 外链保留标签，防止来源绝对路径进入成果。
    text = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", text)
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1])
        else:
            paragraph.add_run(part)


def _style_table(table):
    """为索引和正文数据表设置可读、可跨页的原生 Word 样式。"""
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    for row_number, row in enumerate(table.rows):
        if row_number == 0:
            header = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(header)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            margins = properties.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                properties.append(margins)
            for side in ("top", "left", "bottom", "right"):
                value = margins.find(qn(f"w:{side}"))
                if value is None:
                    value = OxmlElement(f"w:{side}")
                    margins.append(value)
                value.set(qn("w:w"), "100")
                value.set(qn("w:type"), "dxa")
            if row_number == 0 or row_number % 2 == 0:
                shade = properties.find(qn("w:shd"))
                if shade is None:
                    shade = OxmlElement("w:shd")
                    properties.append(shade)
                shade.set(qn("w:fill"), "1F4E78" if row_number == 0 else "F2F6FA")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if row_number == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def _markdown(document, text, skip_initial_heading=None):
    from docx.oxml import OxmlElement
    lines, index = text.splitlines(), 0
    first_content = True
    while index < len(lines):
        line = lines[index].strip()
        if not line or re.fullmatch(r"[-*_]{3,}", line):
            index += 1
            continue
        if line.startswith("```"):
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            document.add_paragraph("\n".join(code))
            index += 1
            continue
        if "|" in line and index + 1 < len(lines) and re.fullmatch(r"\s*\|?\s*:?-{3,}.*", lines[index + 1]) and "|" in lines[index + 1]:
            rows = [[cell.strip() for cell in line.strip("|").split("|")]]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            columns = max(len(row) for row in rows)
            table = document.add_table(rows=0, cols=columns)
            table.style = "Table Grid"
            for number, row in enumerate(rows):
                cells = table.add_row().cells
                for col, value in enumerate(row):
                    _inline(cells[col].paragraphs[0], value)
            _style_table(table)
            first_content = False
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        if heading:
            if not (first_content and skip_initial_heading and heading.group(2).strip() == skip_initial_heading.strip()):
                document.add_heading(heading.group(2), level=min(len(heading.group(1)) + 1, 4))
        elif re.match(r"^[-*+]\s+", line):
            _inline(document.add_paragraph(style="List Bullet"), re.sub(r"^[-*+]\s+", "", line))
        elif re.match(r"^\d+[.)、]\s*", line):
            _inline(document.add_paragraph(style="List Number"), re.sub(r"^\d+[.)、]\s*", "", line))
        else:
            _inline(document.add_paragraph(), line)
        first_content = False
        index += 1


def _document(title, volume):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = document.styles[name]
        style.font.name = "黑体"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        # Word 内置 Title 在部分安装中带主题色边框；正式投标标题只使用字体和留白。
        paragraph_properties = style._element.get_or_add_pPr()
        borders = paragraph_properties.find(qn("w:pBdr"))
        if borders is not None:
            paragraph_properties.remove(borders)
    document.add_paragraph(title, "Title")
    document.add_paragraph(volume)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("第 ")
    _field(footer, " PAGE ", "1")
    footer.add_run(" 页")
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    return document


def _source_images(project, evidence, output_dir):
    from PIL import Image, ImageOps
    files = {item["id"]: item for item in project.load("files", [])}
    file = files.get(evidence.get("file_id"))
    if not file:
        raise ValueError(f"证据 {evidence['id']} 缺少原文件记录")
    source = project.safe_path(file["path"])
    if not source.is_file():
        raise ValueError(f"证据原文件不存在：{file['path']}")
    actual_hash = sha256_file(source)
    recorded = evidence.get("file_sha256") or file.get("sha256")
    if recorded and recorded != actual_hash:
        raise ValueError(f"证据 {evidence['id']} 的原文件已变化，须重新核验")
    result = []
    if source.suffix.lower() == ".pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(source)
        try:
            pages = evidence.get("pages") or []
            if not pages:
                raise ValueError(f"证据 {evidence['id']} 尚未指定需组卷的原始页码")
            for number in dict.fromkeys(pages):
                if not isinstance(number, int) or number < 1 or number > len(pdf):
                    raise ValueError(f"证据 {evidence['id']} 的原始页码超出范围：{number}")
                page = pdf[number - 1]
                bitmap = page.render(scale=200 / 72)
                image_path = output_dir / f"{_bookmark(evidence['id'])}_p{number}.png"
                try:
                    bitmap.to_pil().save(image_path)
                finally:
                    bitmap.close()
                    page.close()
                result.append({"image": image_path, "source_page": number, "file_id": file["id"], "file_sha256": actual_hash})
        finally:
            pdf.close()
    elif source.suffix.lower() in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        with Image.open(source) as image:
            count = getattr(image, "n_frames", 1)
            pages = evidence.get("pages") or list(range(1, count + 1))
            for number in dict.fromkeys(pages):
                if not isinstance(number, int) or number < 1 or number > count:
                    raise ValueError(f"证据 {evidence['id']} 的图片页码超出范围：{number}")
                image.seek(number - 1)
                image_path = output_dir / f"{_bookmark(evidence['id'])}_p{number}.png"
                ImageOps.exif_transpose(image).convert("RGB").save(image_path)
                result.append({"image": image_path, "source_page": number, "file_id": file["id"], "file_sha256": actual_hash})
    else:
        raise ValueError(f"证据附件暂支持 PDF 和图片，请先将 {source.name} 转为保留原貌的 PDF")
    return result


def _assemble_model(project, split):
    settings = project.load("settings", {})
    sections = project.load("sections", [])
    forms = project.load("forms", [])
    evidence = project.load("evidence", [])
    responses = project.load("responses", [])
    # 组卷映射只采用当前规则下通过确定性核验的响应。候选和证明不足项留在
    # 缺件报告中，避免在最终待签章文件里混入未支持材料。
    from .matching import match
    checked = match(project)
    valid_response_ids = {response["id"] for rule in checked["rules"] for response in rule["responses"] if response.get("id") and response["valid"]}
    responses = [response for response in responses if response.get("id") in valid_response_ids]
    selected_ids = {eid for response in responses for eid in response.get("evidence_ids", [])}
    evidence = [item for item in evidence if item["id"] in selected_ids]
    targets, volumes = {}, {}
    def add(item, kind, default):
        volume = item.get("volume") or default if split else "投标文件"
        targets[item["id"]] = {"id": item["id"], "kind": kind, "volume": volume, "title": item.get("title", item["id"]), "bookmark": _bookmark(item["id"]), "record": item}
        volumes.setdefault(volume, []).append(item["id"])
    for item in forms:
        add(item, "form", "商务部分")
    for item in evidence:
        add(item, "evidence", "商务部分")
    for item in sections:
        add(item, "section", "技术部分")
    if not volumes:
        raise ValueError("没有可组卷的正文、表单或响应证据")
    configured = [value if isinstance(value, str) else value.get("name") for value in settings.get("volumes", [])]
    order = [name for name in configured if name in volumes] + [name for name in ("商务部分", "技术部分") if name in volumes and name not in configured]
    order += [name for name in volumes if name not in order]
    rules = [rule for rule in project.load("rules", []) if rule.get("status") != "retired" and rule.get("kind") in {"qualification", "business", "technical"}]
    indexes = []
    for rule in rules:
        ids = []
        for response in responses:
            if response.get("rule_id") == rule["id"] and response.get("status") != "rejected":
                ids += response.get("evidence_ids", []) + response.get("section_ids", [])
                ids += response.get("form_ids", [])
        ids += [section["id"] for section in sections if rule["id"] in section.get("rule_ids", [])]
        indexes.append({"rule_id": rule["id"], "title": rule.get("title", ""), "kind": rule["kind"], "max_score": rule.get("max_score"), "target_ids": list(dict.fromkeys(ids)), "missing_targets": [target for target in ids if target not in targets]})
    return targets, volumes, order, indexes


def _run_word(project, job_path, result_path):
    if os.name != "nt":
        raise RuntimeError("当前组卷需要 Windows 桌面 Word；其他环境只能生成未分页审阅稿")
    script = Path(__file__).parent / "assets" / "word_export.ps1"
    executable = shutil.which("powershell") or shutil.which("pwsh")
    if not executable:
        raise RuntimeError("未找到 PowerShell，无法完成 Word 实际分页")
    try:
        completed = subprocess.run([executable, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script), "-JobPath", str(job_path), "-ResultPath", str(result_path)], capture_output=True, timeout=300, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("Word 分页超过五分钟，未将成果标记完成；未终止任何用户 Word 进程") from exc
    if not result_path.is_file():
        detail = completed.stderr.decode("utf-8", errors="replace")[-800:]
        raise RuntimeError("Word 未返回分页结果：" + detail)
    result = json.loads(result_path.read_text(encoding="utf-8-sig"))
    if completed.returncode or result.get("status") != "ok":
        raise RuntimeError("Word 分页失败：" + result.get("error", "未提供错误详情"))
    return result


def _ready(project):
    assembly = project.load("assembly", {})
    fingerprint = assembly_fingerprint(project)
    errors = _integrity_errors(project)
    if not assembly.get("rendered") or assembly.get("input_fingerprint") != fingerprint:
        errors.append("需先用当前资料生成并分页审阅稿")
    if not _confirmed(project, "assembly", fingerprint):
        errors.append("当前组卷输入尚无用户确认")
    if not _confirmed(project, "visual", visual_fingerprint(project)):
        errors.append("当前 Word/PDF 及页码位置尚无视觉确认")
    from .workflow import is_confirmed
    if not is_confirmed(project, "rules"):
        errors.append("当前招标规则版本尚未确认")
    if project.load("evidence", []) and not is_confirmed(project, "selection"):
        errors.append("当前人员、业绩及证明材料选择尚未确认")
    if project.load("sections", []):
        for scope, title in (("brief", "正式文本沟通记录"), ("plan", "技术策划"), ("draft", "Markdown正文")):
            if not is_confirmed(project, scope):
                errors.append(f"当前{title}版本尚未确认")
    for section in project.load("sections", []):
        if section.get("status") != "confirmed":
            errors.append(f"章节 {section['id']} 未确认或已失效")
    for form in project.load("forms", []):
        if form.get("status") != "confirmed" or not form.get("template_approved") or form.get("errors"):
            errors.append(f"表单 {form['id']} 或其招标模板尚未核准")
    for rule in project.load("rules", []):
        if rule.get("status") == "draft" or rule.get("conflict"):
            errors.append(f"招标规则 {rule['id']} 尚未确认或存在冲突")
    for issue in project.load("issues", []):
        if issue.get("status", "open") == "open" and issue.get("severity") == "error":
            errors.append("终审阻断项：" + issue.get("message", issue.get("id", "")))
    errors += assembly.get("errors", [])
    # 待签章定稿必须先对本程序组卷成品完成人工与Agent的内容双审；
    # 不能只靠 assembly/visual 两个确认绕过新门槛。
    from .final_review import assembly_gate
    gate = assembly_gate(project)
    if not gate["ok"]:
        errors.append("待签章定稿需要当前组卷成品的内容双审通过：" + gate["reason"])
    # 组卷成功只证明版式链路可用。资格、证据、技术复核、表单和一致性
    # 仍须通过当前主记录重新审核，不能靠旧报告或单独的组卷确认绕过。
    from .matching import audit
    audit_result = audit(project)
    blocking_categories = {
        "rules", "scope", "coverage", "evidence", "stale", "section", "review",
        "placeholder", "form", "consistency", "facts", "score", "assembly",
    }
    for issue in audit_result.get("issues", []):
        if issue.get("severity") == "error" or issue.get("category") in blocking_categories:
            errors.append("终审未通过：" + issue.get("message", issue.get("id", "")))
    if errors:
        raise ValueError("不能生成待签章定稿：" + "；".join(dict.fromkeys(errors)))
    verification = verify_output(project)
    if verification["errors"]:
        raise ValueError("输出核验未通过：" + "；".join(verification["errors"]))
    destination = project.safe_path(f"07_最终输出/{assembly['run_id']}")
    destination.mkdir(parents=True, exist_ok=True)
    old_to_new = {}
    for output in assembly["outputs"]:
        for kind in ("docx", "pdf"):
            relative = output.get(kind)
            if not relative:
                continue
            original = project.safe_path(relative)
            copied = destination / original.name
            if copied.exists() and sha256_file(copied) != sha256_file(original):
                raise ValueError(f"最终输出已有不同内容，保留人工文件：{copied.name}")
            if original.resolve() != copied.resolve():
                shutil.copy2(original, copied)
            new_relative = copied.relative_to(project.root).as_posix()
            old_to_new[relative] = new_relative
            output[kind] = new_relative
    positions = project.load("positions", [])
    for position in positions:
        position["output_path"] = old_to_new.get(position.get("output_path"), position.get("output_path"))
    assembly.update({"mode": "ready", "status": "ready_for_signature", "signature_status": "pending", "published_at": utc_now(), "verified": True, "final_review_id": gate["review_id"], "manual_checks": project.load("manual_checks", [])})
    # 待签章成品单独锁存快照；后续build review覆盖assembly记录时，旧ready成品仍可追溯。
    hashes = {}
    for output in assembly["outputs"]:
        hashes[output["volume"]] = {kind: sha256_file(project.safe_path(output[kind])) for kind in ("docx", "pdf") if output.get(kind) and project.safe_path(output[kind]).is_file()}
    history = project.load("ready_history", [])
    snapshot = {"id": "READY-" + str(assembly.get("run_id")), "run_id": assembly.get("run_id"), "at": utc_now(),
                "outputs": [{"volume": output["volume"], "docx": output.get("docx"), "pdf": output.get("pdf")} for output in assembly["outputs"]],
                "hashes": hashes, "positions": positions, "final_review_id": gate["review_id"],
                "signature_status": "pending", "input_fingerprint": fingerprint}
    history = [row for row in history if row.get("run_id") != snapshot["run_id"]] + [snapshot]
    project.commit({"positions": positions, "assembly": assembly, "ready_history": history}, reason="按已确认审阅稿原字节输出待签章定稿")
    return assembly


def build(project, mode="review", split=False, render=True) -> dict:
    if mode not in {"review", "ready"}:
        raise ValueError("组卷模式只能为 review 或 ready")
    settings = project.load("settings", {})
    if settings.get("anonymous"):
        raise ValueError("已识别暗标要求；首版不支持暗标组卷，请人工处理")
    if mode == "ready":
        return _ready(project)
    from docx.shared import Cm
    from PIL import Image
    targets, volumes, order, indexes = _assemble_model(project, split)
    run_id = utc_now().replace(":", "").replace("-", "").replace(".", "") + "_" + uuid4().hex[:6]
    folder = project.safe_path(f"06_审核检查/组卷预览/{run_id}")
    images = project.cache_dir / "assembly" / run_id
    folder.mkdir(parents=True, exist_ok=True)
    images.mkdir(parents=True, exist_ok=True)
    errors, warnings, positions, source_map = [], [], [], []
    jobs, outputs = [], []
    fingerprint = assembly_fingerprint(project)
    for volume_number, volume in enumerate(order, 1):
        document = _document(project.meta.get("name", "投标文件"), volume)
        bookmarks, links, crossrefs, insertions = [], [], [], []
        document.add_paragraph("目录", style="Title")
        _field(document.add_paragraph(), ' TOC \\o "1-3" \\h \\z \\u ', "目录将在 Word 分页后更新")
        if settings.get("allow_indices", True):
            for kind, title in [("qualification", "资格审查响应索引"), ("business", "商务评分响应索引"), ("technical", "技术评分响应索引")]:
                relevant = [row for row in indexes if row["kind"] == kind]
                if not relevant:
                    continue
                document.add_heading(title, level=1)
                table = document.add_table(rows=1, cols=3)
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, ["编号与要求", "响应材料或章节", "册号与页码"]):
                    cell.text = text
                for index_row in relevant:
                    cells = table.add_row().cells
                    cells[0].text = index_row["rule_id"] + " " + index_row["title"]
                    usable = [targets[value] for value in index_row["target_ids"] if value in targets]
                    if not usable:
                        cells[1].text = "待补充响应"
                        cells[2].text = "待定位"
                    for target_number, target in enumerate(usable):
                        label = target["id"] + " " + target["title"]
                        paragraph = cells[1].paragraphs[0] if target_number == 0 else cells[1].add_paragraph()
                        location = cells[2].paragraphs[0] if target_number == 0 else cells[2].add_paragraph()
                        if target["volume"] == volume:
                            _link(paragraph, target["bookmark"], label)
                            location.add_run("第 ")
                            _field(location, f" PAGEREF {target['bookmark']} \\h ")
                            location.add_run(" 页")
                            links.append({"rule_id": index_row["rule_id"], "target_id": target["id"], "bookmark": target["bookmark"], "minimum_links": 2})
                        else:
                            paragraph.add_run(label)
                            variable = "LOC_" + json_hash({"volume": target["volume"], "target": target["id"]})[:16]
                            _field(location, f" DOCVARIABLE {variable} ")
                            crossrefs.append({"variable": variable, "target_id": target["id"], "target_volume": target["volume"]})
                _style_table(table)
        for target_id in volumes[volume]:
            target = targets[target_id]
            record = target["record"]
            document.add_page_break()
            heading = document.add_heading(target["title"], level=1)
            _anchor(heading, target["bookmark"], len(bookmarks) + 1)
            bookmarks.append({"target_id": target_id, "bookmark": target["bookmark"]})
            if target["kind"] == "section":
                path = project.safe_path(record["path"])
                if not path.is_file():
                    errors.append(f"章节文件缺失：{record['path']}")
                    document.add_paragraph("正文缺失，待补充")
                else:
                    if record.get("sha256") and record["sha256"] != sha256_file(path):
                        warnings.append(f"章节 {target_id} 存在人工修改，需重新确认")
                    _markdown(document, path.read_text(encoding="utf-8-sig"), skip_initial_heading=target["title"])
            elif target["kind"] == "form":
                path = project.safe_path(record["output_path"]) if record.get("output_path") else None
                if not path or not path.is_file():
                    errors.append(f"表单 {target_id} 尚未填充")
                    document.add_paragraph("表单尚未填充")
                elif path.suffix.lower() == ".md":
                    _markdown(document, path.read_text(encoding="utf-8-sig"))
                elif path.suffix.lower() == ".docx":
                    insertion = "INS_" + json_hash(target_id)[:16]
                    placeholder = document.add_paragraph("表单内容待插入")
                    _anchor(placeholder, insertion, len(bookmarks) + 10000)
                    insertions.append({"bookmark": insertion, "path": str(path)})
                else:
                    errors.append(f"表单 {target_id} 的输出格式不受支持")
            else:
                try:
                    sources = _source_images(project, record, images)
                    for number, source in enumerate(sources):
                        if number:
                            document.add_page_break()
                        paragraph = document.add_paragraph(f"{target['title']}  来源第 {source['source_page']} 页")
                        page_target = f"{target_id}_p{source['source_page']}"
                        page_bookmark = _bookmark(page_target)
                        _anchor(paragraph, page_bookmark, len(bookmarks) + 1)
                        bookmarks.append({"target_id": page_target, "bookmark": page_bookmark})
                        with Image.open(source["image"]) as picture:
                            width, height = picture.size
                        scale = min(17.0 / width, 21.5 / height)
                        picture_paragraph = document.add_paragraph()
                        picture_paragraph.paragraph_format.space_after = 0
                        picture_paragraph.add_run().add_picture(str(source["image"]), width=Cm(width * scale), height=Cm(height * scale))
                        source_map.append({"evidence_id": target_id, "target_id": page_target, "volume": volume, "file_id": source["file_id"], "source_page": source["source_page"], "file_sha256": source["file_sha256"]})
                except (ValueError, RuntimeError) as exc:
                    errors.append(str(exc))
                    document.add_paragraph("证据附件未能组入，须处理原始材料后重试")
        safe_volume = re.sub(r'[<>:"/\\|?*]', "_", volume)
        docx = folder / f"{volume_number:02d}_{safe_volume}.docx"
        pdf = docx.with_suffix(".pdf")
        document.save(docx)
        outputs.append({"volume": volume, "docx": docx.relative_to(project.root).as_posix(), "pdf": pdf.relative_to(project.root).as_posix(), "bookmarks": bookmarks, "expected_links": links, "crossrefs": crossrefs})
        jobs.append({"volume": volume, "docx": str(docx), "pdf": str(pdf), "bookmarks": bookmarks, "crossrefs": crossrefs, "insertions": insertions})
    for row in indexes:
        if not row["target_ids"] or row["missing_targets"]:
            message = f"{row['rule_id']} 未完整绑定可组卷的响应内容"
            (errors if row["kind"] == "qualification" else warnings).append(message)
    index_lines = ["# 组卷内部响应索引", "", "|要求|内容|册号|", "|---|---|---|"]
    for row in indexes:
        for target_id in row["target_ids"] or ["待补充"]:
            target = targets.get(target_id, {})
            index_lines.append(f"|{row['rule_id']} {row['title']}|{target_id} {target.get('title', '')}|{target.get('volume', '待定位')}|")
    write_text(folder / "内部响应索引.md", "\n".join(index_lines) + "\n")
    result = {"run_id": run_id, "mode": "review", "status": "unpaginated", "input_fingerprint": fingerprint, "outputs": outputs, "rendered": False, "verified": False, "errors": list(dict.fromkeys(errors)), "warnings": list(dict.fromkeys(warnings)), "source_map": source_map, "signature_status": "pending", "split": split, "created_at": utc_now()}
    if render:
        job_path, result_path = images / "word_job.json", images / "word_result.json"
        atomic_json(job_path, {"volumes": jobs})
        try:
            word = _run_word(project, job_path, result_path)
            result["word"] = word
            for output in outputs:
                volume_result = next(value for value in word["volumes"] if value["volume"] == output["volume"])
                output["page_count"] = volume_result["page_count"]
                output["fields"] = volume_result.get("fields", [])
                for position in volume_result["positions"]:
                    positions.append({"id": output["volume"] + ":" + position["target_id"], "volume": output["volume"], **position, "output_path": output["pdf"]})
            result.update({"rendered": True, "status": "review"})
        except RuntimeError as exc:
            result["errors"].append(str(exc))
    for output in outputs:
        output["hashes"] = {kind: sha256_file(project.safe_path(output[kind])) for kind in ("docx", "pdf") if project.safe_path(output[kind]).is_file()}
    project.commit({"positions": positions, "assembly": result}, reason="生成组卷审阅稿，签章仍待人工办理")
    if result["rendered"]:
        verification = verify_output(project)
        result["verification"] = verification
        result["verified"] = not verification["errors"]
        result["errors"] = list(dict.fromkeys(result["errors"] + verification["errors"]))
        project.save("assembly", result, reason="记录实际 PDF 页码与链接核验")
    return result


def _page_number(reader, destination):
    from pypdf.generic import ArrayObject, IndirectObject
    if isinstance(destination, str):
        value = reader.named_destinations.get(destination)
        return reader.get_destination_page_number(value) + 1 if value else None
    if isinstance(destination, IndirectObject):
        destination = destination.get_object()
    if isinstance(destination, ArrayObject) and destination:
        reference = destination[0]
        if isinstance(reference, int):
            return reference + 1
        for index, page in enumerate(reader.pages, 1):
            if page.indirect_reference == reference:
                return index
    return None


def _outline_positions(reader):
    found = {}
    for name, destination in reader.named_destinations.items():
        found[name] = reader.get_destination_page_number(destination) + 1
    def walk(items):
        for item in items:
            if isinstance(item, list):
                walk(item)
            elif getattr(item, "title", None):
                try:
                    found[item.title] = reader.get_destination_page_number(item) + 1
                except Exception:
                    continue
    walk(reader.outline)
    return found


def verify_output(project) -> dict:
    from pypdf import PdfReader
    from lxml import etree
    assembly = project.load("assembly", {})
    errors, warnings, details = [], [], []
    if not assembly.get("rendered") or not assembly.get("outputs"):
        return {"status": "failed", "errors": ["尚未完成 Word 实际分页与 PDF 导出"], "warnings": [], "volumes": []}
    if assembly.get("input_fingerprint") != assembly_fingerprint(project):
        errors.append("组卷输入已经变化，现有输出需重新生成")
    positions = {(row["volume"], row["target_id"]): row for row in project.load("positions", [])}
    for output in assembly.get("outputs", []):
        volume = output["volume"]
        pdf_value, docx_value = output.get("pdf"), output.get("docx")
        if not pdf_value or not docx_value:
            errors.append(f"{volume} 缺少Word或PDF路径记录，无法核验")
            continue
        pdf, docx = project.safe_path(pdf_value), project.safe_path(docx_value)
        if not pdf.is_file() or not docx.is_file():
            errors.append(f"{volume} 的 Word 或 PDF 文件缺失")
            continue
        for kind, path in (("docx", docx), ("pdf", pdf)):
            if output.get("hashes", {}).get(kind) != sha256_file(path):
                errors.append(f"{volume} 的 {kind} 文件已改变，原核验不再有效")
        try:
            reader = PdfReader(pdf)
            if len(reader.pages) != output.get("page_count"):
                errors.append(f"{volume} PDF 页数与 Word 实际分页不一致")
            bookmark_pages = _outline_positions(reader)
            for target in output.get("bookmarks", []):
                position = positions.get((volume, target["target_id"]))
                if not position or not isinstance(position.get("pdf_page"), int) or not 1 <= position["pdf_page"] <= len(reader.pages):
                    errors.append(f"{volume} 目标 {target['target_id']} 缺少有效分页位置")
                elif bookmark_pages.get(target["bookmark"]) != position["pdf_page"]:
                    errors.append(f"{volume} PDF 书签 {target['target_id']} 的实际目的页不一致或缺失")
            destinations = Counter()
            link_count = 0
            for page in reader.pages:
                for annotation_ref in page.get("/Annots", []):
                    annotation = annotation_ref.get_object()
                    if annotation.get("/Subtype") != "/Link":
                        continue
                    action = annotation.get("/A", {})
                    if hasattr(action, "get_object"):
                        action = action.get_object()
                    if action.get("/S") in {"/GoToR", "/Launch"}:
                        errors.append(f"{volume} 含不支持的跨文件或本机链接")
                        continue
                    if action.get("/S") == "/URI":
                        uri = str(action.get("/URI", ""))
                        if uri.startswith(("file:", "\\\\")) or re.match(r"^[A-Za-z]:", uri):
                            errors.append(f"{volume} 含开发机绝对路径链接")
                        continue
                    destination = annotation.get("/Dest") or action.get("/D")
                    if destination is not None:
                        number = _page_number(reader, destination)
                        if number is None or not 1 <= number <= len(reader.pages):
                            errors.append(f"{volume} 存在无效 PDF 内部链接目的地")
                        else:
                            destinations[number] += 1
                            link_count += 1
            expected = Counter()
            for link in output.get("expected_links", []):
                position = positions.get((volume, link["target_id"]))
                if position:
                    expected[position["pdf_page"]] += link.get("minimum_links", 1)
            for page, count in expected.items():
                if destinations[page] < count:
                    errors.append(f"{volume} 第 {page} 页缺少预期的索引跳转链接")
            pageref_count = 0
            for field in output.get("fields", []):
                match = re.search(r"\bPAGEREF\s+(\S+)", field.get("code", ""), re.I)
                if match:
                    pageref_count += 1
                    target = next((row for row in positions.values() if row["volume"] == volume and row["bookmark"] == match.group(1)), None)
                    if target is None and match.group(1).startswith("_Toc"):
                        target = field
                    if not target or str(field.get("result", "")).strip() != str(target.get("printed_page")):
                        errors.append(f"{volume} PAGEREF 页码与实际目标页不一致")
                    if target and destinations.get(target.get("pdf_page"), 0) < 1:
                        errors.append(f"{volume} 目录或索引页码引用缺少对应 PDF 跳转")
            for crossref in output.get("crossrefs", []):
                target = positions.get((crossref["target_volume"], crossref["target_id"]))
                expected_value = f"{crossref['target_volume']} 第 {target['printed_page']} 页" if target else None
                fields = [field for field in output.get("fields", []) if re.search(r"\bDOCVARIABLE\s+" + re.escape(crossref["variable"]) + r"\b", field.get("code", ""))]
                if not target or not fields or any(field.get("result", "").strip() != expected_value for field in fields):
                    errors.append(f"{volume} 跨册索引与目标册实际页码不一致")
            with ZipFile(docx) as archive:
                for filename in archive.namelist():
                    if filename.endswith(".rels"):
                        root = etree.fromstring(archive.read(filename))
                        for relation in root:
                            value = relation.get("Target", "")
                            if relation.get("TargetMode") == "External" and (value.startswith(("file:", "\\\\")) or re.match(r"^[A-Za-z]:", value)):
                                errors.append(f"{volume} Word 含本机绝对路径链接")
            details.append({"volume": volume, "page_count": len(reader.pages), "bookmarks_checked": len(output.get("bookmarks", [])), "internal_links": link_count, "pageref_checked": pageref_count})
        except Exception as exc:
            errors.append(f"{volume} 输出读取失败：{exc}")
    return {"status": "passed" if not errors else "failed", "errors": list(dict.fromkeys(errors)), "warnings": warnings, "volumes": details}
