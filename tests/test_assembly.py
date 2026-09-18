"""组卷测试使用合成项目；Word 集成验证须显式设置 BIDFLOW_WORD_TEST=1。"""
from __future__ import annotations

import os
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image, ImageDraw

from bidflow.assembly import assembly_fingerprint, build, verify_output, visual_fingerprint
from bidflow.forms import fill_form, price_upper
from bidflow.project import Project
from bidflow.utils import json_hash, sha256_file, write_text
from bidflow.workflow import confirm
from final_review_helpers import complete_content


def make_project(root):
    project = Project.create(root, "合成测试：咨询服务投标")
    text = "# 工作组织\n软件测试合成正文，不可用于真实投标。\n\n|序号|工作任务|成果|\n|---|---|---|\n"
    text += "\n".join(f"|{number}|核对测试资料并形成相应工作记录|测试成果{number}|" for number in range(1, 46))
    first = project.safe_path("05_投标文件编制/SEC001.md")
    second = project.safe_path("05_投标文件编制/SEC002.md")
    write_text(first, text)
    write_text(second, "# 质量保障\n按资料核对、校核、复核三个步骤组织测试。\n")
    image_path = project.safe_path("01_输入文件/02_公司证照/测试证照.png")
    image = Image.new("RGB", (1500, 800), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 30, 1470, 770), outline="black", width=4)
    draw.text((80, 80), "SYNTHETIC TEST EVIDENCE - LANDSCAPE", fill="black", font_size=32)
    image.save(image_path)
    tender_path = project.safe_path("01_输入文件/01_招标文件/合成招标要求.md")
    write_text(tender_path, "法人资格。证照完整。工作方案。质量保障。")
    template_path = project.safe_path("01_输入文件/08_其他补充资料/测试投标函.docx")
    template = Document()
    template.add_heading("合成测试投标函", 0)
    line = template.add_paragraph("项目：")
    line.add_run("{{project_").bold = True
    line.add_run("name}}")
    table = template.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "投标人"
    table.cell(0, 1).text = "{{company_name}}"
    table.cell(1, 0).text = "报价（元）"
    table.cell(1, 1).text = "{{price}} / {{price_upper}}"
    template.sections[0].header.paragraphs[0].text = "测试模板：{{company_name}}"
    template.save(template_path)
    project.commit({
        "facts": {"project_name": project.meta["name"], "company_name": "合成测试公司", "deadline": "2026-12-31", "service_period": "60日", "price": "10020.05", "price_upper": "壹万零贰拾元零伍分", "formal_brief": {"questions": ["文种用途", "发文主体", "主要读者", "决策目标", "正文范围"], "document_type": "测试技术方案", "purpose": "软件组卷测试", "author": "合成测试公司", "audience": "测试评委", "decision_goal": "验证响应覆盖", "body_scope": "合成测试正文", "sources": ["合成招标要求"], "structure": ["工作方案", "质量保障"], "length": "测试长度", "tone": "正式直接"}},
        "sections": [
            {"id": "SEC001", "title": "工作方案", "path": first.relative_to(project.root).as_posix(), "sha256": sha256_file(first), "rule_ids": ["T001"], "status": "confirmed", "writer_id": "writer"},
            {"id": "SEC002", "title": "质量保障", "path": second.relative_to(project.root).as_posix(), "sha256": sha256_file(second), "rule_ids": ["T002"], "status": "confirmed", "writer_id": "writer"},
        ],
        "rules": [
            {"id": "Q001", "title": "法人资格", "kind": "qualification", "status": "confirmed", "revision": 1, "sources": [{"file_id": "DOC_TENDER", "block_id": "B_TENDER", "quote": "法人资格", "page": 1}]},
            {"id": "S001", "title": "证照完整", "kind": "business", "status": "confirmed", "revision": 1, "max_score": 5, "criteria": {"scoring": {"method": "fixed", "points": 5}}, "sources": [{"file_id": "DOC_TENDER", "block_id": "B_TENDER", "quote": "证照完整", "page": 1}]},
            {"id": "T001", "title": "工作方案", "kind": "technical", "status": "confirmed", "revision": 1, "max_score": 10, "subrequirements": ["步骤完整"], "sources": [{"file_id": "DOC_TENDER", "block_id": "B_TENDER", "quote": "工作方案", "page": 1}]},
            {"id": "T002", "title": "质量保障", "kind": "technical", "status": "confirmed", "revision": 1, "max_score": 5, "subrequirements": ["三级复核"], "sources": [{"file_id": "DOC_TENDER", "block_id": "B_TENDER", "quote": "质量保障", "page": 1}]},
        ],
        "files": [{"id": "DOC_TENDER", "path": tender_path.relative_to(project.root).as_posix(), "sha256": sha256_file(tender_path), "category": "招标文件", "page_count": 1}, {"id": "DOC_TEST", "path": image_path.relative_to(project.root).as_posix(), "sha256": sha256_file(image_path), "category": "公司证照", "page_count": 1}],
        "blocks": [{"id": "B_TENDER", "file_id": "DOC_TENDER", "text": "法人资格。证照完整。工作方案。质量保障。", "page": 1, "quality": "ok"}, {"id": "B_EVIDENCE", "file_id": "DOC_TEST", "text": "合成测试证照", "page": 1, "quality": "ok"}],
        "analysis_coverage": [{"id": "B_TENDER", "block_id": "B_TENDER", "block_hash": json_hash({"id": "B_TENDER", "file_id": "DOC_TENDER", "text": "法人资格。证照完整。工作方案。质量保障。", "page": 1, "quality": "ok"}), "task_id": "TEST", "actor": "测试规则分析"}],
        "evidence": [{"id": "E001", "title": "合成测试证照（横版）", "file_id": "DOC_TEST", "file_sha256": sha256_file(image_path), "pages": [1], "verification": "verified", "verified_by": "evidence-reviewer", "facts": {"license": "合成测试"}}],
        "responses": [
            {"id": "RESP001", "rule_id": "Q001", "evidence_ids": ["E001"], "form_ids": ["FORM001"], "status": "supported", "rationale": "证照及投标函共同响应", "rule_revision": 1},
            {"id": "RESP002", "rule_id": "S001", "evidence_ids": ["E001"], "status": "supported", "rationale": "已核验测试证照", "rule_revision": 1},
            {"id": "RESP003", "rule_id": "T001", "section_ids": ["SEC001"], "status": "supported", "rationale": "工作方案章节响应", "rule_revision": 1, "covered_subrequirements": ["步骤完整"]},
            {"id": "RESP004", "rule_id": "T002", "section_ids": ["SEC002"], "status": "supported", "rationale": "质量保障章节响应", "rule_revision": 1, "covered_subrequirements": ["三级复核"]},
        ],
        "forms": [{"id": "FORM001", "title": "合成测试投标函", "template_path": template_path.relative_to(project.root).as_posix(), "template_approved": True, "template_sha256": sha256_file(template_path), "required_fields": ["project_name", "company_name", "price", "price_upper"], "source_refs": [{"file_id": "DOC_TENDER", "block_id": "B_TENDER"}]}],
        "plans": [{"id": "SEC001", "title": "工作方案", "rule_ids": ["T001"], "content_points": ["步骤完整"], "status": "draft"}, {"id": "SEC002", "title": "质量保障", "rule_ids": ["T002"], "content_points": ["三级复核"], "status": "draft"}],
    })
    filled = fill_form(project, "FORM001")
    assert not filled["errors"]
    evidence_block = next(row for row in project.load("blocks") if row["id"] == "B_EVIDENCE")
    project.save("evidence_coverage", [{"id": "ECOV001", "task_id": "TEST_EVIDENCE", "block_id": "B_EVIDENCE", "block_hash": json_hash(evidence_block), "file_id": "DOC_TEST", "file_sha256": sha256_file(image_path)}], reason="合成测试：登记证据覆盖")
    forms = project.load("forms")
    forms[0]["status"] = "confirmed"
    project.save("forms", forms)
    current_facts = project.fingerprint(["facts"])
    project.save("reviews", [
        {"id": "REV001", "section_id": "SEC001", "writer_id": "writer", "reviewer_id": "reviewer", "section_sha256": sha256_file(first), "rule_revisions": {"T001": 1}, "facts_fingerprint": current_facts, "round": 1, "status": "pass", "scores": [{"rule_id": "T001", "score": 8, "reason": "测试评分"}], "covered_subrequirements": {"T001": ["步骤完整"]}, "findings": []},
        {"id": "REV002", "section_id": "SEC002", "writer_id": "writer", "reviewer_id": "reviewer", "section_sha256": sha256_file(second), "rule_revisions": {"T002": 1}, "facts_fingerprint": current_facts, "round": 1, "status": "pass", "scores": [{"rule_id": "T002", "score": 4, "reason": "测试评分"}], "covered_subrequirements": {"T002": ["三级复核"]}, "findings": []},
    ])
    confirm(project, "rules", "测试规则确认", attest_human=True)
    confirm(project, "selection", "测试材料确认", attest_human=True)
    confirm(project, "brief", "测试正文沟通确认", attest_human=True)
    confirm(project, "plan", "测试策划确认", attest_human=True)
    confirm(project, "draft", "测试正文确认", attest_human=True)
    return project


@pytest.mark.parametrize("amount,expected", [(0, "零元整"), (0.05, "零元伍分"), (10, "壹拾元整"), (10020.05, "壹万零贰拾元零伍分"), (100000001, "壹亿零壹元整"), (100010001, "壹亿零壹万零壹元整")])
def test_price_upper(amount, expected):
    assert price_upper(amount) == expected


@pytest.mark.parametrize("amount", ["NaN", "Infinity", "-1", "不是数字"])
def test_bad_price(amount):
    with pytest.raises(ValueError):
        price_upper(amount)


def test_editable_form_and_source_preserved(tmp_path):
    project = make_project(tmp_path / "project")
    form = project.load("forms")[0]
    original = sha256_file(project.safe_path(form["template_path"]))
    with ZipFile(project.safe_path(form["output_path"])) as archive:
        content = archive.read("word/document.xml").decode()
        header = archive.read("word/header1.xml").decode()
    assert "{{" not in content + header
    assert "合成测试公司" in content and "合成测试公司" in header
    assert "<w:tbl>" in content and "<w:b" in content
    old = form["output_path"]
    fill_form(project, "FORM001")
    assert project.safe_path(old).is_file()
    assert original == sha256_file(project.safe_path(form["template_path"]))


def test_form_missing_and_shared_conflict(tmp_path):
    project = make_project(tmp_path / "project")
    forms = project.load("forms")
    forms[0]["fields"] = {"company_name": "错误公司"}
    forms[0]["required_fields"].append("contact")
    project.save("forms", forms)
    result = fill_form(project, "FORM001")
    assert "contact" in result["missing_fields"]
    assert any("共用事实" in message for message in result["errors"])


def test_unpaginated_and_anonymous_are_not_final(tmp_path):
    project = make_project(tmp_path / "project")
    result = build(project, render=False)
    assert not result["rendered"] and not result["verified"]
    with pytest.raises(ValueError, match="分页审阅稿"):
        build(project, mode="ready")
    project.save("settings", {"anonymous": True})
    with pytest.raises(ValueError, match="暗标"):
        build(project)


def test_candidate_evidence_is_not_inserted_into_assembly(tmp_path):
    project = make_project(tmp_path / "project")
    evidence = project.load("evidence")
    evidence[0]["verification"] = "candidate"
    evidence[0]["verified_by"] = ""
    project.save("evidence", evidence)
    result = build(project, render=False)
    assert not result["source_map"]
    assert all("E001" not in {item["target_id"] for item in output.get("bookmarks", [])} for output in result["outputs"])


@pytest.mark.skipif(os.environ.get("BIDFLOW_WORD_TEST") != "1", reason="需要显式运行 Windows Word 集成测试")
def test_real_word_split_repaginate_and_ready(tmp_path):
    project = make_project(tmp_path / "word_project")
    first = build(project, split=True)
    assert first["rendered"] and first["verified"], first["errors"]
    assert len(first["outputs"]) == 2
    assert first["source_map"][0]["source_page"] == 1
    assert all(volume["internal_links"] >= 2 for volume in first["verification"]["volumes"])
    before = {row["target_id"]: row["pdf_page"] for row in project.load("positions")}
    sections = project.load("sections")
    first_path = project.safe_path(sections[0]["path"])
    write_text(first_path, first_path.read_text(encoding="utf-8") + "\n\n" + "\n\n".join("追加软件分页测试段落。" * 10 for _ in range(40)))
    assert verify_output(project)["errors"]
    sections[0]["sha256"] = sha256_file(first_path)
    project.save("sections", sections)
    reviews = project.load("reviews")
    reviews[0]["section_sha256"] = sections[0]["sha256"]
    reviews[0]["id"] = "REV003"
    reviews[0]["round"] = 2
    project.save("reviews", reviews)
    confirm(project, "draft", "测试正文再次确认", attest_human=True)
    second = build(project, split=True)
    assert second["verified"], second["errors"]
    after = {row["target_id"]: row["pdf_page"] for row in project.load("positions")}
    assert after["SEC002"] > before["SEC002"]
    confirm(project, "assembly", "合成测试确认者", attest_human=True)
    confirm(project, "visual", "合成测试确认者", attest_human=True)
    # 待签章门禁要求当前组卷成品先完成人工与Agent的内容双审。
    content_review = complete_content(project, writer="writer", agent="成品双审Agent", human="成品人工复核人")
    hashes = [dict(output["hashes"]) for output in second["outputs"]]
    published = build(project, mode="ready")
    assert published["status"] == "ready_for_signature" and published["signature_status"] == "pending"
    assert published["final_review_id"] == content_review
    assert [output["hashes"] for output in published["outputs"]] == hashes
    assert verify_output(project)["status"] == "passed"
    assert all(output["docx"].startswith("07_最终输出/") for output in published["outputs"])
    assert build(project, mode="ready")["status"] == "ready_for_signature"
    project.safe_path(published["outputs"][0]["pdf"]).write_bytes(b"changed")
    assert verify_output(project)["errors"]
    with pytest.raises(ValueError):
        build(project, mode="ready")
