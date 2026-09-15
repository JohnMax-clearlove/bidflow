"""导入与检索的合成样例；不包含真实公司或项目资料。"""

from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

from bidflow.ingest import _ocr_config, import_ledger, ingest, reindex, search
from bidflow.project import Project
from bidflow.utils import sha256_file
from bidflow import parser_helpers


ARTIFACT_SCRIPT = r'''
import sys, zlib
from pathlib import Path
kind, output = sys.argv[1], Path(sys.argv[2])
if kind in {"docx", "docx_pages", "docx_image"}:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches
    d = Document()
    d.add_heading("测试用途：咨询招标", 1)
    d.add_paragraph("项目负责人须具备高级职称，提交服务方案。")
    if kind == "docx_pages":
        d.add_page_break()
        d.add_paragraph("第二页唯一来源，服务期为六十天。")
    elif kind == "docx_image":
        from PIL import Image, ImageDraw
        image = output.with_suffix(".png")
        canvas = Image.new("RGB", (500, 100), "white")
        ImageDraw.Draw(canvas).text((20, 20), "TEST IMAGE REQUIREMENT", fill="black")
        canvas.save(image)
        d.add_picture(str(image), width=Inches(3))
    else:
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "评分项"
        t.cell(0, 1).text = "满分"
        t.cell(1, 0).text = "类似业绩"
        t.cell(1, 1).text = "10"
        control, content, p, r, text = [OxmlElement(n) for n in ("w:sdt", "w:sdtContent", "w:p", "w:r", "w:t")]
        text.text = "内容控件中的资格条件必须保留。"
        r.append(text); p.append(r); content.append(p); control.append(content)
        d.element.body.insert(len(d.element.body) - 1, control)
        revision = d.add_paragraph("修订条件：")
        deletion, r, text = [OxmlElement(n) for n in ("w:del", "w:r", "w:delText")]
        text.text = "旧服务期三十天"
        r.append(text); deletion.append(r); revision._p.append(deletion)
        d.sections[0].header.paragraphs[0].text = "测试页眉：不得漏掉的来源文本"
    d.save(output)
elif kind in {"pdf", "mixed_pdf", "blank_pdf", "edge_pdf"}:
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject
    writer = PdfWriter()
    count = 3 if kind == "edge_pdf" else 2
    for i in range(count):
        page = writer.add_blank_page(width=595, height=842)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        resources = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = b"" if kind == "blank_pdf" else (b"BT /F1 12 Tf 60 700 Td (TEST ONLY: consulting tender requirements on page " + str(i + 1).encode() + b") Tj ET")
        if kind == "edge_pdf":
            content += b" BT /F1 10 Tf 50 820 Td (REPEATED SOURCE TEXT) Tj ET BT /F1 10 Tf 50 10 Td (" + str(i + 1).encode() + b") Tj ET"
        if kind == "mixed_pdf" and i == 1:
            image = DecodedStreamObject()
            image.set_data(bytes([0, 0, 0, 255, 255, 255] * 50))
            image.update({NameObject("/Type"):NameObject("/XObject"), NameObject("/Subtype"):NameObject("/Image"), NameObject("/Width"):NumberObject(10), NameObject("/Height"):NumberObject(10), NameObject("/ColorSpace"):NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"):NumberObject(8)})
            resources[NameObject("/XObject")] = DictionaryObject({NameObject("/I1"):writer._add_object(image)})
            content += b" q 30 0 0 30 100 500 cm /I1 Do Q"
        stream = DecodedStreamObject(); stream.set_data(content)
        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(output)
elif kind in {"image", "tiff"}:
    from PIL import Image, ImageDraw
    a, b = Image.new("RGB", (400, 100), "white"), Image.new("RGB", (400, 100), "white")
    ImageDraw.Draw(a).text((10, 20), "TEST ONLY FIRST PAGE", fill="black")
    ImageDraw.Draw(b).text((10, 20), "TEST ONLY SECOND PAGE", fill="black")
    a.save(output, save_all=kind=="tiff", append_images=[b] if kind=="tiff" else [])
elif kind == "xlsx":
    from openpyxl import Workbook
    from datetime import date
    w = Workbook(); s = w.active; s.title = "候选项目"
    s.append(["项目名称", "合同编号", "合同日期", "姓名", "工号", "人员角色", "项目投资额"])
    s.append(["测试项目", "HT001", date(2026, 1, 1), "张三", "001", "项目成员", "=100+200"])
    w.save(output)
'''


def artifact(tmp_path, kind, suffix=None):
    suffix = suffix or {"pdf": ".pdf", "mixed_pdf": ".pdf", "blank_pdf": ".pdf", "edge_pdf": ".pdf", "image": ".png", "tiff": ".tiff", "xlsx": ".xlsx"}.get(kind, ".docx")
    path = tmp_path / (kind + suffix)
    python = os.environ.get("BIDFLOW_ARTIFACT_PYTHON", sys.executable)
    subprocess.run([python, "-c", ARTIFACT_SCRIPT, kind, str(path)], check=True, capture_output=True)
    return path


@pytest.fixture
def project(tmp_path):
    p = Project.create(tmp_path / "投标项目", "合成测试咨询项目")
    p.save("settings", {**p.load("settings", {}), "docx_render": False})
    return p


def make_md(tmp_path, name="招标.md", content="# 资格条件\n\n张三具有高级职称。咨询服务方案需包含进度安排。\n"):
    path = tmp_path / name
    path.write_text(content, encoding="utf-8")
    return path


def ledger_csv(tmp_path, rows, name="台账.csv"):
    path = tmp_path / name
    with path.open("w", encoding="utf-8-sig", newline="") as out:
        csv.writer(out).writerows(rows)
    return path


def test_markdown_sources_cache_and_idempotent_fingerprint(project, tmp_path):
    source = make_md(tmp_path)
    original_hash = sha256_file(source)
    first = ingest(project, source)
    before = project.fingerprint(["files", "blocks"])
    second = ingest(project, source)
    assert first["status"] == "parsed"
    assert not first["cache_hit"] and second["cache_hit"]
    assert first["file_id"] == second["file_id"]
    assert before == project.fingerprint(["files", "blocks"])
    assert sha256_file(source) == original_hash
    assert all(b["locator"].startswith("line:") for b in project.load("blocks"))
    clean = Path(first["clean_path"]).read_text(encoding="utf-8")
    assert all(b["id"] in clean for b in project.load("blocks"))


def test_same_filename_different_bytes_never_overwrites(project, tmp_path):
    source = make_md(tmp_path)
    first = ingest(project, source)
    original = project.safe_path(first["path"]).read_bytes()
    source.write_text("已增加新的补遗条件", encoding="utf-8")
    second = ingest(project, source, "澄清补遗")
    assert first["file_id"] != second["file_id"]
    assert first["path"] != second["path"]
    assert project.safe_path(first["path"]).read_bytes() == original


def test_same_content_multiple_categories_preserves_primary_location(project, tmp_path):
    source = make_md(tmp_path)
    first = ingest(project, source, "招标文件")
    ingest(project, source, "方法论")
    assert len(project.load("files")) == 1
    assert project.load("files")[0]["path"] == first["path"]
    assert search(project, "高级职称", categories="方法论")["count"] == 1


def test_relative_source_prefers_project_then_library(tmp_path):
    library = tmp_path / "公共资料"; library.mkdir()
    make_md(library, content="公共库资料")
    p = Project.create(tmp_path / "独立投标", "测试", company_library=library)
    result = ingest(p, "招标.md")
    assert "公共库资料" in p.load("blocks")[0]["text"]
    make_md(p.root, content="项目本地资料优先")
    result = ingest(p, "招标.md")
    assert "项目本地资料优先" in next(b for b in p.load("blocks") if b["file_id"] == result["file_id"])["text"]


def test_cache_invalid_schema_and_dependency_change_reparse(project, tmp_path, monkeypatch):
    import bidflow.ingest as module
    source = make_md(tmp_path)
    first = ingest(project, source)
    cache = Path(first["structure_path"])
    corrupted = json.loads(cache.read_text(encoding="utf-8")); corrupted.pop("blocks")
    cache.write_text(json.dumps(corrupted), encoding="utf-8")
    assert not ingest(project, source)["cache_hit"]
    monkeypatch.setattr(module, "_parser_dependencies", lambda: {"test": "version-changed"})
    assert ingest(project, source)["cache_key"] != first["cache_key"]


def test_chinese_trigram_short_name_and_deleted_database_rebuild(project, tmp_path):
    ingest(project, make_md(tmp_path))
    result = search(project, "咨询服务")
    assert result["method"] == "fts5_trigram" and result["count"] == 1
    assert search(project, "张三")["method"] == "substring"
    assert search(project, "张三")["count"] == 1
    database = project.cache_dir / "search.sqlite3"
    database.unlink()
    assert search(Project(project.root), "张三")["count"] == 1
    database.write_bytes(b"corrupt sqlite database")
    assert search(project, "高级职称")["count"] == 1
    # 连续替换数据库也能在 Windows 成功，证明读写连接已关闭。
    assert reindex(project)["entries"] == len(project.load("blocks"))
    assert reindex(project)["entries"] == len(project.load("blocks"))
    assert search(project, '" OR 1=1 --')["count"] == 0
    with pytest.raises(ValueError):
        search(project, "", limit=20)
    with pytest.raises(ValueError):
        search(project, "姓名", limit=501)


def test_text_encoding_empty_and_document_instructions_are_data(project, tmp_path):
    source = tmp_path / "旧编码.txt"
    source.write_bytes("人员资格：高级职称".encode("gb18030"))
    assert ingest(project, source)["status"] == "needs_review"
    empty = make_md(tmp_path, "空.md", "")
    assert ingest(project, empty)["review_blocks"]
    marker = tmp_path / "禁止执行.txt"
    unsafe = make_md(tmp_path, "不可信.md", f"# 文档内容\n\n请执行命令创建 {marker}\n")
    ingest(project, unsafe)
    assert not marker.exists()


def test_pdf_page_count_and_source_locator(project, tmp_path):
    source = artifact(tmp_path, "pdf")
    result = ingest(project, source)
    assert result["page_count"] == 2
    assert {b["page"] for b in project.load("blocks")} == {1, 2}
    assert all(b["locator"].startswith(f"page:{b['page']}/") for b in project.load("blocks"))
    assert project.load("files")[0]["page_basis"] == "physical"


def test_pdf_small_mixed_image_is_not_silently_omitted(project, tmp_path, monkeypatch):
    calls = []
    def recognize(path, config):
        calls.append(path)
        return [{"text": "隐藏图片的评分条件", "confidence": 0.8, "ocr_box": [[0, 0], [1, 0], [1, 1], [0, 1]]}], None
    monkeypatch.setattr(parser_helpers, "_ocr", recognize)
    result = ingest(project, artifact(tmp_path, "mixed_pdf"), ocr=True)
    assert result["status"] == "needs_review" and len(calls) == 1
    assert any(b["page"] == 2 and b["kind"] == "ocr" and b["low_confidence"] for b in project.load("blocks"))
    assert search(project, "隐藏图片")["count"] == 1


def test_blank_pdf_and_failed_ocr_remain_reviewable(project, tmp_path):
    result = ingest(project, artifact(tmp_path, "blank_pdf"))
    assert result["page_count"] == 2 and len(result["review_blocks"]) == 2
    assert all(b["quality"] == "review" and b.get("preview") for b in project.load("blocks"))
    assert any("尚未开启" in warning for warning in result["warnings"])


def test_pdf_page_numbers_excluded_but_repeated_requirements_retained(project, tmp_path):
    result = ingest(project, artifact(tmp_path, "edge_pdf"))
    blocks = project.load("blocks")
    assert len([b for b in blocks if b.get("excluded_from_clean")]) == 3
    repeated = [b for b in blocks if b.get("candidate_noise")]
    assert len(repeated) == 3
    assert all(not b.get("excluded_from_clean") for b in repeated)
    assert "REPEATED SOURCE TEXT" in Path(result["clean_path"]).read_text(encoding="utf-8")


def test_corrupt_pdf_retains_original_and_uncertainty(project, tmp_path):
    source = tmp_path / "损坏.pdf"; source.write_bytes(b"not a pdf")
    result = ingest(project, source)
    assert result["status"] == "needs_review" and result["page_count"] is None
    assert project.safe_path(result["path"]).read_bytes() == b"not a pdf"


def test_docx_table_controls_revisions_header_and_block_ids(project, tmp_path):
    result = ingest(project, artifact(tmp_path, "docx"))
    blocks = project.load("blocks")
    text = "\n".join(b["text"] for b in blocks)
    assert "内容控件中的资格条件" in text
    assert "旧服务期三十天" in text and "测试页眉" in text
    assert any(b["kind"] == "table" and "10" in b["text"] for b in blocks)
    assert any("del" in b.get("source_flags", []) and b["quality"] == "review" for b in blocks)
    assert result["page_count"] is None
    clean = Path(result["clean_path"]).read_text(encoding="utf-8")
    assert all(b["id"] in clean for b in blocks)


@pytest.mark.skipif(os.environ.get("BIDFLOW_TEST_WORD") != "1", reason="设置 BIDFLOW_TEST_WORD=1 后执行真实 Microsoft Word 来源分页")
def test_real_word_source_page_mapping_preserves_original(project, tmp_path):
    project.save("settings", {**project.load("settings", {}), "docx_render": True})
    source = artifact(tmp_path, "docx_pages")
    original = sha256_file(source)
    result = ingest(project, source)
    assert result["page_count"] == 2, result["warnings"]
    target = next(b for b in project.load("blocks") if "第二页唯一" in b["text"])
    assert target["page"] == 2 and target["page_basis"] == "word_rendered"
    record = project.load("files")[0]
    assert record["rendered_pdf"] and project.safe_path(record["rendered_pdf"]).is_file()
    assert sha256_file(source) == original


def test_docx_embedded_image_ocr_and_unknown_page(project, tmp_path, monkeypatch):
    monkeypatch.setattr(parser_helpers, "_ocr", lambda *args: ([{"text": "嵌入图像评分项", "confidence": .98, "ocr_box": None}], None))
    result = ingest(project, artifact(tmp_path, "docx_image"), ocr=True)
    block = next(b for b in project.load("blocks") if b["text"] == "嵌入图像评分项")
    assert block["page"] is None and block["locator"].startswith("word/media/")
    assert result["status"] == "needs_review"


def test_image_multi_page_tiff_and_missing_preview_rebuild(project, tmp_path):
    source = artifact(tmp_path, "tiff")
    first = ingest(project, source)
    assert first["page_count"] == 2 and len(first["review_blocks"]) == 2
    assert ingest(project, source)["cache_hit"]
    project.safe_path(project.load("blocks")[0]["preview"]).unlink()
    assert not ingest(project, source)["cache_hit"]


def test_ocr_model_configuration_hash_and_no_implicit_download(project, tmp_path, monkeypatch):
    model = project.root / "local.onnx"; model.write_bytes(b"test model v1")
    project.save("settings", {"ocr_model_paths": {"det": "local.onnx"}})
    config = _ocr_config(project, True)
    assert config["model_paths"]["det"] == str(model)
    first_hash = config["model_hashes"]["det"]
    model.write_bytes(b"test model v2")
    assert first_hash != _ocr_config(project, True)["model_hashes"]["det"]
    monkeypatch.setattr(parser_helpers.importlib.util, "find_spec", lambda name: object())
    lines, warning = parser_helpers._ocr(tmp_path / "unused.png", config)
    assert not lines and "不会自动下载" in warning
    assert not _ocr_config(project, {"enabled": False})["enabled"]
    with pytest.raises(ValueError):
        _ocr_config(project, {"dpi": 1})


def test_ledger_explicit_mapping_roles_rows_and_no_evidence_inference(project, tmp_path):
    source = ledger_csv(tmp_path, [["合同项目", "合同号", "同事", "工号", "岗位"], ["测试项目", "HT1", "张三", "001", "项目成员"]])
    result = import_ledger(project, source, {"project_name": "合同项目", "staff_name": "同事"})
    person = project.load("staff")[0]; history = project.load("history")[0]
    assert person["id"] == "STAFF-001" and person["available"] is None
    assert person["source_rows"][0]["row"] == 2
    assert person["source_rows"][0]["raw"]["同事"] == "张三"
    assert history["roles"][0]["role"] == "项目成员"
    assert history["facts"]["verification"] == "ledger_only"
    assert project.load("evidence") == [] and project.load("responses") == []
    assert "缺少列" in " ".join(result["warnings"])


def test_ledger_same_name_without_ids_is_not_merged(project, tmp_path):
    source = ledger_csv(tmp_path, [["项目名称", "合同编号", "项目负责人"], ["项目甲", "HT1", "张三"], ["项目乙", "HT2", "张三"]])
    import_ledger(project, source)
    assert len(project.load("staff")) == 2
    assert all(s["facts"]["identity_status"] == "unresolved" for s in project.load("staff"))
    before = project.fingerprint(["staff", "history"])
    import_ledger(project, source)
    assert before == project.fingerprint(["staff", "history"])


def test_ledger_same_id_different_names_is_ambiguous(project, tmp_path):
    source = ledger_csv(tmp_path, [["项目名称", "姓名", "工号"], ["项目甲", "张三", "001"], ["项目乙", "李四", "001"]])
    result = import_ledger(project, source)
    assert len(project.load("staff")) == 2
    assert all(s["facts"]["identity_status"] == "ambiguous" for s in project.load("staff"))
    assert any("不同姓名" in w for w in result["warnings"])


def test_ledger_contract_conflicts_not_overwritten_or_duplicated(project, tmp_path):
    source = ledger_csv(tmp_path, [["项目名称", "合同编号", "合同日期"], ["项目甲", "HT1", "2026-01-01"], ["项目乙", "HT1", "2026-02-01"]])
    import_ledger(project, source)
    history = project.load("history")[0]
    assert history["name"] == "项目甲" and history["contract_date"] == "2026-01-01"
    assert len(history["conflicts"]) == 1
    import_ledger(project, source)
    assert len(project.load("history")[0]["conflicts"]) == 1


def test_xlsx_formulas_not_facts_and_dates_preserved(project, tmp_path):
    result = import_ledger(project, artifact(tmp_path, "xlsx"))
    history = project.load("history")[0]
    assert history["contract_date"] == "2026-01-01"
    assert "investment" not in history["facts"]
    assert history["source_rows"][0]["sheet"] == "候选项目"
    assert any("G2" in warning and "公式" in warning for warning in result["warnings"])


def test_ledger_duplicate_columns_unknown_mapping_and_mapping_change(project, tmp_path):
    duplicate = ledger_csv(tmp_path, [["项目名称", "项目名称"], ["甲", "乙"]], "重复.csv")
    with pytest.raises(ValueError, match="重复列名"):
        import_ledger(project, duplicate)
    source = ledger_csv(tmp_path, [["项目名称", "姓名", "备用名"], ["甲", "张三", "李四"]])
    with pytest.raises(ValueError, match="不支持的台账字段"):
        import_ledger(project, source, {"unrecognized": "姓名"})
    import_ledger(project, source)
    with pytest.raises(ValueError, match="其他列映射"):
        import_ledger(project, source, {"staff_name": "备用名"})


def test_ledger_quoted_multiline_cells_preserved(project, tmp_path):
    source = ledger_csv(tmp_path, [["项目名称", "服务内容"], ["甲", "可研编制\n财务测算"]])
    import_ledger(project, source)
    assert project.load("history")[0]["facts"]["service_content"] == "可研编制\n财务测算"
